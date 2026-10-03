"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (không gõ tay).

Tên cột sheet "Experiments" (giữ nguyên): exp_id, group, description, loss, optimizer, lr, weight_decay,
batch, epochs, hidden, dropout, clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch,
final_train_loss, final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
eval_acc, eval_macro_f1, figure_file, notes. Các cột công thức ở cuối bảng mẫu tự tính, không ghi đè.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

LOSS_LABEL = {"ce": "CE", "mse": "MSE"}
OPT_LABEL = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
FORMULA_COLS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
MAX_ROWS = 60   # số dòng có sẵn công thức trong mẫu


def _clean(o):
    """Chuyển numpy/torch scalar về kiểu Python để json.dump được."""
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if hasattr(o, "item") and not isinstance(o, (str, bytes)):
        return o.item()
    return o


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    with open(path, "w") as f:
        json.dump(_clean({k: result[k] for k in ("cfg", "history", "summary")}), f, indent=1)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi <results_dir>/*.json là kết quả thí nghiệm (có khoá cfg), sắp theo exp_id."""
    out = []
    for p in Path(results_dir).glob("*.json"):
        with open(p) as f:
            r = json.load(f)
        if "cfg" in r and "history" in r:
            r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
            out.append(r)
    return sorted(out, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một dòng của bảng: cfg + summary (+ eval_acc/eval_macro_f1 nếu có) + figure_file."""
    c, s = result["cfg"], result["summary"]
    row = dict(
        exp_id=c["exp_id"], group=c.get("group", "other"), description=c.get("description", ""),
        loss=LOSS_LABEL[c["loss"]], optimizer=OPT_LABEL[c["optimizer"]], lr=c["lr"],
        weight_decay=c.get("weight_decay", 0.0), batch=c["batch"], epochs=c["epochs"],
        hidden="-".join(map(str, c["hidden"])), dropout=c.get("dropout", 0.0),
        clip_norm=c["clip_norm"] if c.get("clip_norm") else "none", precision=c["precision"],
        init=c["init"], seed=c["seed"],
        diverged="Y" if s.get("diverged") else "N", figure_file=f"figures/{c['exp_id']}.png", notes=notes,
    )
    for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
              "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        row[k] = s.get(k)
    if eval_scores:
        row["eval_acc"], row["eval_macro_f1"] = eval_scores["accuracy"], eval_scores["macro_f1"]
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" (từ dòng 2), seeds vào sheet "Seeds" (cột A, dòng 2..6)
    và nhận xét nhóm vào cột H của sheet "Summary", rồi lưu thành out_path.
    Không dùng data_only=True (sẽ mất công thức); các cột công thức không bị ghi đè."""
    import openpyxl
    assert len(rows) <= MAX_ROWS, f"mẫu chỉ có sẵn công thức cho {MAX_ROWS} dòng"
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    cols = {ws.cell(1, j).value: j for j in range(1, ws.max_column + 1)}
    data_cols = [c for c in cols if c not in FORMULA_COLS]
    for i in range(MAX_ROWS):                         # xoá giá trị mẫu cũ (vd. dòng baseline điền sẵn)
        for c in data_cols:
            ws.cell(2 + i, cols[c]).value = None
    for i, row in enumerate(rows):
        for c in data_cols:
            v = row.get(c)
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                v = None
            ws.cell(2 + i, cols[c]).value = v
    if seed_ids is not None:
        ss = wb["Seeds"]
        for i in range(5):
            ss.cell(2 + i, 1).value = seed_ids[i] if i < len(seed_ids) else None
    if summary_notes:
        sm = wb["Summary"]
        for r in range(2, 12):
            g = sm.cell(r, 1).value
            if g in summary_notes:
                sm.cell(r, 8).value = summary_notes[g]
    wb.save(out_path)
